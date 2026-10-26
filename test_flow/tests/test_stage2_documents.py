import json
import shutil
import subprocess
from pathlib import Path
from types import SimpleNamespace

import pytest

import kordoc_parser as kp

APP_DIR = Path(__file__).resolve().parent.parent
KORDOC_MODULES = APP_DIR / "kordoc" / "node_modules"


def kordoc_installed() -> bool:
    try:
        kp.kordoc_command()
        return KORDOC_MODULES.exists() and shutil.which("node") is not None
    except kp.KordocUnavailable:
        return False


needs_kordoc = pytest.mark.skipif(not kordoc_installed(), reason="kordoc not installed (cd test_flow/kordoc && npm ci)")

SAMPLE_MD = """# 2026 사업계획

## 1장 개요

본 계획은 회의 자동화 서비스의 2026년 운영 방향과 일정, 담당 조직을 정리한 문서이다.

## 2장 예산

### 2.1 마케팅 예산

올해 마케팅 예산 총액은 10억 원이며 분기별로 나누어 집행한다. 집행 결과는 매 분기 보고한다.

| 분기 | 금액 |
| --- | --- |
| 1분기 | 2억 |
| 2분기 | 3억 |

### 2.2 보안 점검

보안 점검은 분기마다 실시하고 결과를 정보보호위원회에 보고한다. 담당자는 보안팀장이다.
"""


@pytest.fixture(scope="module")
def sample_dir(tmp_path_factory):
    if not kordoc_installed():
        pytest.skip("kordoc not installed")
    out = tmp_path_factory.mktemp("docs")
    script = (
        "const {markdownToHwpx}=require('kordoc');const fs=require('fs');"
        "markdownToHwpx(fs.readFileSync(0,'utf8')).then(b=>fs.writeFileSync(process.argv[1],Buffer.from(b)));"
    )
    subprocess.run(["node", "-e", script, str(out / "사업계획.hwpx")], input=SAMPLE_MD, text=True, check=True, cwd=APP_DIR / "kordoc")

    import fitz

    doc = fitz.open()
    for text in ["1장 개요\n회의 자동화 서비스의 운영 방향을 정리한다. 회의록 생성과 태스크 추출을 자동화한다.",
                 "2장 예산\n올해 마케팅 예산 총액은 10억 원이다. 배포일은 11월 3일로 결정했다. 보안 점검은 분기마다 한다."]:
        page = doc.new_page()
        page.insert_text((72, 90), text, fontname="korea", fontsize=12)
    doc.save(out / "운영안.pdf")

    from docx import Document

    word = Document()
    word.add_heading("인사 규정", level=1)
    word.add_paragraph("연차 휴가는 입사일 기준으로 15일을 부여하며 사용 계획은 매년 1월에 제출한다.")
    word.save(out / "인사규정.docx")

    (out / "메모.txt").write_text("첫 문단은 회의 준비 체크리스트입니다. 자료 공유와 참석자 확인을 포함합니다.\n\n두 번째 문단은 후속 조치입니다.", encoding="utf-8")
    return out


def test_table_markdown_handles_merged_cells():
    table = {
        "rows": 2,
        "cols": 3,
        "cells": [
            [{"text": "구분", "rowSpan": 2}, {"text": "상반기", "colSpan": 2}],
            [{"text": "1분기"}, {"text": "2|분기"}],
        ],
    }
    assert kp.table_markdown(table).splitlines() == [
        "| 구분 | 상반기 | 상반기 |",
        "| --- | --- | --- |",
        "| 구분 | 1분기 | 2\\|분기 |",
    ]


def test_blocks_from_kordoc_builds_heading_path():
    result = {
        "blocks": [
            {"type": "heading", "level": 1, "text": "A", "pageNumber": 1},
            {"type": "heading", "level": 2, "text": "A-1", "pageNumber": 1},
            {"type": "paragraph", "text": "본문1", "pageNumber": 1},
            {"type": "image", "text": "img.png", "pageNumber": 1},
            {"type": "heading", "level": 2, "text": "A-2", "pageNumber": 2},
            {"type": "table", "pageNumber": 2, "table": {"rows": 1, "cols": 1, "cells": [[{"text": "x"}]]}},
        ]
    }
    rows = kp.blocks_from_kordoc(result, doc_id="d", source_path=Path("a.hwpx"), source_sha256="s")
    assert [r["heading_path"] for r in rows] == [["A"], ["A", "A-1"], ["A", "A-1"], ["A", "A-2"], ["A", "A-2"]]
    assert rows[-1]["block_type"] == "table" and rows[-1]["page_no"] == 2
    assert all(r["parser"] == "kordoc" for r in rows)


def test_luna_fallback_only_for_pages_without_text(monkeypatch, tmp_path):
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(b"%PDF")
    result = {"pageQuality": [{"page": 1, "needsOcr": True}, {"page": 2, "needsOcr": True}, {"page": 3, "needsOcr": False}]}
    rows = [{"page_no": 2, "text": "kordoc OCR이 이미 충분히 읽은 두 번째 페이지 텍스트입니다."}]
    calls = []
    monkeypatch.setattr(kp, "_render_pdf_page", lambda path, page: f"png{page}".encode())
    monkeypatch.setattr(kp, "_luna_ocr", lambda data, name, ctype: calls.append(data) or "루나 OCR 결과")
    out = kp.luna_fallback_rows(pdf, result, rows, doc_id="d", sha="s")
    assert calls == [b"png1"]
    assert out[0]["page_no"] == 1 and out[0]["parser"] == "openai_ocr_fallback"

    monkeypatch.setenv("KORDOC_LUNA_FALLBACK", "false")
    assert kp.luna_fallback_rows(pdf, result, rows, doc_id="d", sha="s") == []


def test_missing_kordoc_raises_clear_error(monkeypatch):
    monkeypatch.setenv("KORDOC_BIN", "/nonexistent/kordoc")
    monkeypatch.setattr(kp, "APP_DIR", Path("/nonexistent"))
    monkeypatch.setattr(kp.shutil, "which", lambda _name: None)
    with pytest.raises(kp.KordocUnavailable, match="npm ci"):
        kp.kordoc_command()


def test_ocr_file_bytes_falls_back_when_kordoc_missing(monkeypatch):
    import model_runtime as mr

    monkeypatch.setattr(mr, "OCR_BACKEND", "kordoc")
    # TXT는 kordoc을 거치지 않고 바로 디코딩한다.
    assert mr.ocr_file_bytes("가나다".encode(), "a.txt") == "가나다"

    def unavailable(*a, **k):
        raise kp.KordocUnavailable("missing")

    monkeypatch.setattr(kp, "extract_text", unavailable)
    monkeypatch.setattr(mr, "_openai_ocr_file_bytes", lambda b, n, c="": "openai path")
    assert mr.ocr_file_bytes(b"x", "a.pdf") == "openai path"


@needs_kordoc
def test_kordoc_parses_hwpx_pdf_docx_txt(sample_dir):
    by_name = {}
    for path in sorted(sample_dir.iterdir()):
        rows, record = kp.parse_file(path)
        assert record["status"] == "success", record
        by_name[path.suffix] = rows

    hwpx = by_name[".hwpx"]
    budget = next(r for r in hwpx if "10억" in r["text"])
    assert budget["heading_path"] == ["2026 사업계획", "2장 예산", "2.1 마케팅 예산"]
    table = next(r for r in hwpx if r["block_type"] == "table")
    assert "| 1분기 | 2억 |" in table["text"]

    pdf = by_name[".pdf"]
    assert {r["page_no"] for r in pdf} == {1, 2}
    assert any("배포일은 11월 3일" in r["text"] for r in pdf if r["page_no"] == 2)

    assert any("연차 휴가" in r["text"] for r in by_name[".docx"])
    assert len(by_name[".txt"]) == 2


@needs_kordoc
def test_parser_and_chunker_end_to_end(sample_dir, tmp_path):
    from document_ingest import read_chunks_jsonl, run_chunker, run_parser

    parsed, chunks_dir = tmp_path / "parsed", tmp_path / "chunks"
    info = run_parser(sample_dir, parsed)
    assert info["mode"] == "kordoc" and info["failed"] == 0 and info["pdf_count"] == 4
    chunker = run_chunker(parsed, sample_dir, chunks_dir)
    chunks = read_chunks_jsonl(Path(chunker["output"]))
    assert chunks

    budget = next(c for c in chunks if "10억" in c["document"] and "사업계획" in c["metadata"]["source_filename"])
    assert "[문서: 사업계획.hwpx]" in budget["document"]
    assert "[위치: 2026 사업계획 > 2장 예산 > 2.1 마케팅 예산]" in budget["document"]
    assert budget["metadata"]["heading_path"][-1] == "2.1 마케팅 예산"

    deploy = next(c for c in chunks if "배포일" in c["document"])
    assert deploy["metadata"]["page_start"] == 2 and "[페이지: 2]" in deploy["document"]

    docx = next(c for c in chunks if "연차 휴가" in c["document"])
    assert docx["metadata"]["source_filename"] == "인사규정.docx"
    # 부모-자식 구조: 모든 청크에 부모 텍스트가 있다(② 단순 분할 제거).
    assert all(c.get("parent_text") and c.get("parent_id") for c in chunks)


@needs_kordoc
def test_extract_text_for_ocr_endpoint(sample_dir):
    text = kp.extract_text((sample_dir / "운영안.pdf").read_bytes(), "운영안.pdf")
    assert "[page 1]" in text and "[page 2]" in text and "배포일" in text


def test_meeting_chunks_have_header():
    from meeting_ingest import make_chunks

    req = SimpleNamespace(title="스프린트 회의", meeting_datetime="2026-10-20 10:00", meeting_id="12", project_id="3")
    result = {"summary": "배포일을 11월 3일로 정했다.", "todo_list": [{"title": "배포 준비", "owner": "김규호"}]}
    chunks = make_chunks(req, result)
    assert chunks[0]["document"].startswith("[회의: 스프린트 회의 / 2026-10-20 10:00]\n[항목: 요약]")
    assert any("[항목: 할 일]" in c["document"] and "김규호" in c["document"] for c in chunks)


def test_uploads_all_go_to_parser(tmp_path):
    from internal_docs.service import UploadedDocument, _prepare_uploads

    files = [UploadedDocument("a.pdf", b"%PDF"), UploadedDocument("b.docx", b"PK"), UploadedDocument("c.txt", "가".encode())]
    prepared = _prepare_uploads(files, tmp_path / "raw", tmp_path / "parse", {})
    assert sorted(p.suffix for p in (tmp_path / "parse").iterdir()) == [".docx", ".pdf", ".txt"]
    assert "text_chunks" not in prepared
    assert {m["source_file_extension"] for m in prepared["metadata_by_source_filename"].values()} == {".pdf", ".docx", ".txt"}
